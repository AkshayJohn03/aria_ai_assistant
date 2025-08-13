import os
import io
import fitz
import openpyxl
import docx
import pytesseract
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

FILES_DIR = os.path.join("data", "files")
os.makedirs(FILES_DIR, exist_ok=True)

def create_pdf(text: str, filename: str) -> str:
    path = os.path.join(FILES_DIR, filename)
    c = canvas.Canvas(path, pagesize=A4)
    width, height = A4
    y = height - 72
    for line in text.splitlines():
        c.drawString(72, y, line[:110])
        y -= 18
        if y < 72:
            c.showPage()
            y = height - 72
    c.save()
    return path

def annotate_pdf(src_path: str, notes: str, out_name: str) -> str:
    doc = fitz.open(src_path)
    for page in doc:
        rect = fitz.Rect(50, 50, 550, 120)
        page.insert_textbox(rect, notes, fontsize=10)
    out_path = os.path.join(FILES_DIR, out_name)
    doc.save(out_path)
    doc.close()
    return out_path

def create_or_update_excel(filename: str, sheet: str, cell: str, value: str) -> str:
    path = os.path.join(FILES_DIR, filename)
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(title=sheet)
    ws[cell] = value
    wb.save(path)
    return path

def create_docx(text: str, filename: str) -> str:
    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    path = os.path.join(FILES_DIR, filename)
    d.save(path)
    return path

def ocr_image(image_path: str) -> str:
    img = Image.open(image_path)
    return pytesseract.image_to_string(img) 